from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def apply_cell_styles(
    cell: Any,
    font: dict[str, Any] | None = None,
    fill: dict[str, Any] | None = None,
    alignment: dict[str, Any] | None = None,
) -> None:
    if font:
        cell.font = Font(**font)
    if fill:
        cell.fill = PatternFill(
            fill_type="solid", fgColor=fill.get("color", "FFFFFF")
        )
    if alignment:
        cell.alignment = Alignment(**alignment)


def create_openpyxl_table(ws: Any, ref_address: str, table_name: str) -> Table:
    tab = Table(displayName=table_name, ref=ref_address)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True
    )
    ws.add_table(tab)
    return tab
