from typing import Any

from openpyxl.utils import get_column_letter

from excelpilot.bridge.models import BridgeError, RangeRef, SheetInfo


def list_openpyxl_sheets(wb: Any) -> list[SheetInfo]:
    infos: list[SheetInfo] = []
    for idx, name in enumerate(wb.sheetnames):
        ws = wb[name]
        infos.append(
            SheetInfo(
                name=name,
                index=idx,
                visibility="visible" if ws.sheet_state == "visible" else "hidden",
                row_count=ws.max_row or 0,
                column_count=ws.max_column or 0,
                has_tables=len(ws.tables) > 0,
                has_pivots=False,
                has_charts=False,
            )
        )
    return infos


def get_sheet_used_range(wb: Any, sheet: str) -> RangeRef:
    if sheet not in wb.sheetnames:
        raise BridgeError("SheetNotFound", f"Sheet {sheet} not found")
    ws = wb[sheet]
    col_letter = get_column_letter(ws.max_column or 1)
    return RangeRef(sheet=sheet, address=f"A1:{col_letter}{ws.max_row or 1}")


def add_openpyxl_sheet(wb: Any, name: str | None = None) -> SheetInfo:
    ws = wb.create_sheet(title=name)
    sheets = list_openpyxl_sheets(wb)
    return [s for s in sheets if s.name == ws.title][0]


def rename_openpyxl_sheet(wb: Any, old_name: str, new_name: str) -> SheetInfo:
    if old_name not in wb.sheetnames:
        raise BridgeError("SheetNotFound", f"Sheet {old_name} not found")
    wb[old_name].title = new_name
    sheets = list_openpyxl_sheets(wb)
    return [s for s in sheets if s.name == new_name][0]
